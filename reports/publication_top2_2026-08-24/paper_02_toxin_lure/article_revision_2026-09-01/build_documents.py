from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
MANUSCRIPT_MD = HERE / "manuscript_article.md"
COVER_MD = HERE / "cover_letter_article.md"
MANUSCRIPT_DOCX = HERE / "Injun_Artificial_Ecology_Value_Learning_Article.docx"
COVER_DOCX = HERE / "Injun_Artificial_Ecology_Cover_Letter.docx"
FONT = "Times New Roman"

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
LIST_ITEM = re.compile(r"^(\d+)\.\s+(.*)$")


def font_run(run, size: float = 12, bold=None, italic=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), FONT)
    rpr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, end])
    font_run(run, 10)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_document(doc: Document, double_spaced: bool = True) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(12)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0 if double_spaced else 6)
    normal.paragraph_format.line_spacing = 2.0 if double_spaced else 1.05

    for name, before, after in (("Heading 1", 12, 4), ("Heading 2", 10, 2), ("Heading 3", 8, 2)):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 2.0 if double_spaced else 1.05
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_number(footer)


def add_inline(paragraph, text: str, size: float = 12) -> None:
    position = 0
    for match in INLINE.finditer(text):
        if match.start() > position:
            font_run(paragraph.add_run(text[position : match.start()]), size)
        token = match.group(0)
        if token.startswith("**"):
            font_run(paragraph.add_run(token[2:-2]), size, bold=True)
        elif token.startswith("*"):
            font_run(paragraph.add_run(token[1:-1]), size, italic=True)
        else:
            font_run(paragraph.add_run(token[1:-1]), size)
        position = match.end()
    if position < len(text):
        font_run(paragraph.add_run(text[position:]), size)


def strip_inline(text: str) -> str:
    return text.replace("**", "").replace("`", "").replace("*", "")


def parse_blocks(lines: list[str]):
    blocks = []
    paragraph_buffer: list[str] = []
    index = 0

    def flush_paragraph() -> None:
        if paragraph_buffer:
            blocks.append(("paragraph", " ".join(item.strip() for item in paragraph_buffer)))
            paragraph_buffer.clear()

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("| ") and index + 1 < len(lines) and re.match(r"^\|\s*:?-+", lines[index + 1].strip()):
            flush_paragraph()
            table_lines = [stripped]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
            blocks.append(("table", rows))
            continue
        if stripped.startswith("!["):
            flush_paragraph()
            match = re.search(r"\(([^)]+)\)", stripped)
            blocks.append(("figure", match.group(1)))
        elif stripped.startswith("#### "):
            flush_paragraph()
            blocks.append(("heading4", stripped[5:]))
        elif stripped.startswith("### "):
            flush_paragraph()
            blocks.append(("heading3", stripped[4:]))
        elif stripped.startswith("## "):
            flush_paragraph()
            blocks.append(("heading2", stripped[3:]))
        elif stripped.startswith("# "):
            flush_paragraph()
            blocks.append(("title", stripped[2:]))
        elif LIST_ITEM.match(stripped):
            flush_paragraph()
            match = LIST_ITEM.match(stripped)
            blocks.append(("numbered", (match.group(1), match.group(2))))
        else:
            paragraph_buffer.append(stripped)
        index += 1
    flush_paragraph()
    return blocks


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.cell(r_index, c_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_index == 0:
                set_cell_shading(cell, "E8EEF5")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_index == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            size = 8.0 if len(row) >= 5 else 8.5
            font_run(paragraph.add_run(strip_inline(value)), size, bold=r_index == 0)
    following = doc.add_paragraph()
    following.paragraph_format.space_after = Pt(3)


def add_figure(doc: Document, path: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(HERE / path), width=Inches(6.35))
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        if "stationary" in path:
            doc_pr[0].set("title", "Figure 1. Stationary food-value learning")
            doc_pr[0].set("descr", "Paired seed totals and low-value meals across three time windows.")
        else:
            doc_pr[0].set("title", "Figure 2. Temporal toxicity lure")
            doc_pr[0].set("descr", "Analytic mixed value and finite-horizon lure rates for four learners.")


def build_manuscript() -> None:
    lines = MANUSCRIPT_MD.read_text(encoding="utf-8").splitlines()
    title = lines[0][2:]
    abstract_index = lines.index("## Abstract")
    intro_index = lines.index("## 1. Introduction")
    abstract_lines = lines[abstract_index + 1 : intro_index]
    abstract_text = " ".join(
        line.strip() for line in abstract_lines if line.strip() and not line.startswith("**Keywords:**")
    )
    keywords = next(line.replace("**Keywords:**", "").strip() for line in abstract_lines if line.startswith("**Keywords:**"))

    doc = Document()
    configure_document(doc, double_spaced=True)
    doc.core_properties.title = title
    doc.core_properties.author = "Chisanupong Injun"
    doc.core_properties.subject = "Article for Artificial Life"
    doc.core_properties.keywords = keywords

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_after = Pt(12)
    title_paragraph.paragraph_format.line_spacing = 1.15
    font_run(title_paragraph.add_run(title), 15, bold=True)

    for value in (
        "Chisanupong Injun",
        "Deebuk Phangnga Wittayayon School, Phang Nga 82000, Thailand",
        "ORCID: https://orcid.org/0009-0000-2979-1916",
        "Corresponding author: Chisanupong Injun; chisanupong.injun@gmail.com",
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.paragraph_format.space_after = Pt(1)
        font_run(paragraph.add_run(value), 11)

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(2)
    heading.paragraph_format.line_spacing = 1.0
    font_run(heading.add_run("Abstract"), 12, bold=True)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.space_after = Pt(4)
    add_inline(paragraph, abstract_text, 10.5)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.0
    font_run(paragraph.add_run("Keywords: "), 10.5, bold=True)
    font_run(paragraph.add_run(keywords), 10.5)
    doc.add_page_break()

    in_references = False
    for kind, content in parse_blocks(lines[intro_index:]):
        if kind.startswith("heading"):
            level = {"heading2": 1, "heading3": 2, "heading4": 3}[kind]
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, content)
            in_references = content == "References"
        elif kind == "paragraph":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.widow_control = True
            is_caption = content.startswith("**Figure ")
            is_equation = content.startswith("`") and content.endswith("`")
            is_admin = in_references or content.startswith("Persistent URL:")
            if in_references:
                paragraph.paragraph_format.left_indent = Inches(0.5)
                paragraph.paragraph_format.first_line_indent = Inches(-0.5)
                paragraph.paragraph_format.line_spacing = 2.0
            elif not is_caption and not is_equation:
                paragraph.paragraph_format.first_line_indent = Inches(0.5)
            if is_caption:
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(8)
                paragraph.paragraph_format.keep_with_next = False
            if is_equation:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Inches(0)
            size = 9.5 if is_caption else (11 if in_references else 12)
            add_inline(paragraph, content, size)
        elif kind == "numbered":
            number, text = content
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.first_line_indent = Inches(-0.25)
            font_run(paragraph.add_run(f"{number}. "), 12)
            add_inline(paragraph, text, 12)
        elif kind == "table":
            add_table(doc, content)
        elif kind == "figure":
            add_figure(doc, content)

    doc.save(MANUSCRIPT_DOCX)


def build_cover() -> None:
    lines = COVER_MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc, double_spaced=False)
    doc.sections[0].footer.paragraphs[0].clear()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(7)
    doc.core_properties.title = "Cover letter for Artificial Life"
    doc.core_properties.author = "Chisanupong Injun"

    for kind, content in parse_blocks(lines[1:]):
        if kind == "paragraph":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.widow_control = True
            add_inline(paragraph, content, 11)
        elif kind == "title":
            continue

    doc.save(COVER_DOCX)


if __name__ == "__main__":
    build_manuscript()
    build_cover()
    print(MANUSCRIPT_DOCX)
    print(COVER_DOCX)
