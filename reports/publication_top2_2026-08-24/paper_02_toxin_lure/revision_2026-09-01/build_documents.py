from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


HERE = Path(__file__).resolve().parent
MANUSCRIPT_MD = HERE / "manuscript_revised.md"
COVER_MD = HERE / "cover_letter_revised.md"
MANUSCRIPT_DOCX = HERE / "Injun_Temporal_Toxicity_Letter_Revised.docx"
COVER_DOCX = HERE / "Injun_Cover_Letter_Revised.docx"
FONT = "Times New Roman"


def font_run(run, size=12, bold=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), FONT)
    rpr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])
    font_run(run, 10)


def configure(doc, double_spaced=True):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(12)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0 if double_spaced else 7)
    normal.paragraph_format.line_spacing = 2.0 if double_spaced else 1.08

    for name, before, after in (("Heading 1", 12, 6), ("Heading 2", 10, 4)):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 2.0 if double_spaced else 1.08
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_number(footer)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(paragraph, text, size=12):
    pos = 0
    for match in INLINE.finditer(text):
        if match.start() > pos:
            font_run(paragraph.add_run(text[pos:match.start()]), size)
        token = match.group(0)
        if token.startswith("**"):
            font_run(paragraph.add_run(token[2:-2]), size, bold=True)
        elif token.startswith("*"):
            font_run(paragraph.add_run(token[1:-1]), size, italic=True)
        else:
            font_run(paragraph.add_run(token[1:-1]), size)
        pos = match.end()
    if pos < len(text):
        font_run(paragraph.add_run(text[pos:]), size)


def blocks(lines):
    result, buffer = [], []

    def flush():
        if buffer:
            result.append(("p", " ".join(x.strip() for x in buffer)))
            buffer.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()
        elif stripped.startswith("!["):
            flush()
            result.append(("figure", re.search(r"\(([^)]+)\)", stripped).group(1)))
        elif stripped.startswith("### "):
            flush()
            result.append(("h2", stripped[4:]))
        elif stripped.startswith("## "):
            flush()
            result.append(("h1", stripped[3:]))
        elif stripped.startswith("# "):
            flush()
            result.append(("title", stripped[2:]))
        else:
            buffer.append(stripped)
    flush()
    return result


def build_manuscript():
    lines = MANUSCRIPT_MD.read_text(encoding="utf-8").splitlines()
    title = lines[0][2:]
    doc = Document()
    configure(doc, double_spaced=True)
    doc.core_properties.title = title
    doc.core_properties.author = "Chisanupong Injun"
    doc.core_properties.subject = "Letter for Artificial Life"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    font_run(p.add_run(title), 14, bold=True)
    for value in (
        "Chisanupong Injun",
        "Deebuk Phangnga Wittayayon School, Phang Nga 82000, Thailand",
        "ORCID: https://orcid.org/0009-0000-2979-1916",
        "Corresponding author: Chisanupong Injun; chisanupong.injun@gmail.com",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(2)
        font_run(p.add_run(value), 11)

    abstract_idx = lines.index("## Abstract")
    intro_idx = lines.index("## 1. Introduction")
    abstract_lines = lines[abstract_idx + 1:intro_idx]
    abstract_text = " ".join(x.strip() for x in abstract_lines if x.strip() and not x.startswith("**Keywords:**"))
    keywords = next(x.replace("**Keywords:**", "").strip() for x in abstract_lines if x.startswith("**Keywords:**"))
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    font_run(p.add_run("Abstract"), 12, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    add_inline(p, abstract_text)
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    font_run(p.add_run("Keywords: "), 11, bold=True)
    font_run(p.add_run(keywords), 11)
    doc.add_page_break()

    in_refs = False
    for kind, text in blocks(lines[intro_idx:]):
        if kind in ("h1", "h2"):
            p = doc.add_paragraph(style="Heading 1" if kind == "h1" else "Heading 2")
            add_inline(p, text)
            in_refs = text == "References"
        elif kind == "figure":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run()
            run.add_picture(str(HERE / text), width=Inches(6.3))
            doc_pr = run._r.xpath(".//wp:docPr")
            if doc_pr:
                doc_pr[0].set("title", "Figure 1. Intermittent safety creates a lure")
                doc_pr[0].set("descr", "Analytic mixture ranking and lure rates for four learning rules.")
        elif kind == "p":
            p = doc.add_paragraph()
            p.paragraph_format.widow_control = True
            is_caption = text.startswith("**Figure ")
            if in_refs:
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.5)
            elif not is_caption and not text.startswith("`Vmix"):
                p.paragraph_format.first_line_indent = Inches(0.5)
            if is_caption:
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(8)
            if text.startswith("`Vmix"):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, text, 10 if is_caption else 12)

    doc.save(MANUSCRIPT_DOCX)


def build_cover():
    lines = COVER_MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure(doc, double_spaced=False)
    doc.sections[0].footer.paragraphs[0].clear()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(5)
    doc.core_properties.title = "Cover letter for Artificial Life"
    doc.core_properties.author = "Chisanupong Injun"

    paragraphs, buffer = [], []
    signature = False
    for line in lines[2:]:
        stripped = line.strip()
        if stripped == "Sincerely,":
            if buffer:
                paragraphs.append(("body", " ".join(buffer)))
                buffer.clear()
            paragraphs.append(("signature", stripped))
            signature = True
        elif signature and stripped:
            paragraphs.append(("signature", stripped))
        elif not stripped:
            if buffer:
                paragraphs.append(("body", " ".join(buffer)))
                buffer.clear()
        else:
            buffer.append(stripped)
    if buffer:
        paragraphs.append(("body", " ".join(buffer)))

    for index, (kind, text) in enumerate(paragraphs):
        p = doc.add_paragraph()
        if index == 0:
            p.paragraph_format.space_after = Pt(12)
        if kind == "signature":
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
        add_inline(p, text, 11)
    doc.save(COVER_DOCX)


if __name__ == "__main__":
    build_manuscript()
    build_cover()
    print(MANUSCRIPT_DOCX)
    print(COVER_DOCX)
